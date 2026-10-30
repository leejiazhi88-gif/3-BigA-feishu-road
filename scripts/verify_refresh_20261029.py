#!/usr/bin/env python3
"""Verify the exported 2026-10-29 refresh against the copied template."""
from __future__ import annotations

import json
import math
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TARGET_FILE = ROOT / 'outputs' / 'refresh_20261029' / '2026-10-29-all.xlsx'
SOURCE_FILE = ROOT / 'outputs' / 'annual-change-color-fix' / '2026-10-23-all.xlsx'
OUTPUT = ROOT / 'exports' / 'refresh_20261029' / 'final_verification.json'
TARGET_SHEET = '2026-10-29（全部）'
SOURCE_SHEET = '2026-10-23（全部）'


def color(cell):
    value = cell.font.color
    if value is None or value.type != 'rgb':
        return '000000'
    return str(value.rgb)[-6:].upper()


def fill(cell):
    value = cell.fill.fgColor
    return str(value.rgb)[-6:].upper() if value.type == 'rgb' else ''


def number(value):
    if value in (None, ''):
        return None
    match = re.match(r'[-+]?\d+(?:\.\d+)?', str(value))
    return float(match.group()) if match else None


def main():
    target = openpyxl.load_workbook(TARGET_FILE, data_only=False)[TARGET_SHEET]
    source = openpyxl.load_workbook(SOURCE_FILE, data_only=False)[SOURCE_SHEET]
    headers = {cell.value: cell.column for cell in target[1]}
    source_headers = {cell.value: cell.column for cell in source[1]}
    target_by_code = {
        str(target.cell(row, headers['代码']).value): row
        for row in range(2, target.max_row + 1) if target.cell(row, 1).value
    }
    source_by_code = {
        str(source.cell(row, source_headers['代码']).value): row
        for row in range(2, source.max_row + 1) if source.cell(row, 1).value
    }

    note_mismatches = []
    composition_mismatches = []
    for code, source_row in source_by_code.items():
        target_row = target_by_code[code]
        for header in ('笔记（AI别刷新这部分，我自己记录）', '笔记2 决策（AI也别刷新）'):
            old = source.cell(source_row, source_headers[header])
            new = target.cell(target_row, headers[header])
            if (old.value, fill(old), color(old)) != (new.value, fill(new), color(new)):
                note_mismatches.append((code, header))
        for header in ('25年收入构成', '25年利润构成', '26年收入构成', '26年利润构成'):
            if source.cell(source_row, source_headers[header]).value != target.cell(
                    target_row, headers[header]).value:
                composition_mismatches.append((code, header))

    mapping = json.loads((ROOT / 'config' / 'industry_color_map_20261023.json').read_text())
    industry_fill_errors = []
    font_rule_errors = []
    pe_headers = ('PE(TTM)', '2026 PE(预测)', '2027 PE(预测)', '2028 PE(预测)')
    yoy_headers = ('2025净利润同比%(实际/预测)', '2026净利润同比%(预测)',
                   '2027净利润同比%(预测)', '2028净利润同比%(预测)')
    actual_headers = (
        '2026Q1净利润(亿元)', '2026Q2锁定预测(亿元)', '2026Q2净利润(亿元)',
        '2026 Q3净利润(亿元) 锁定比较基础', '2026Q3净利润(亿元)',
        '2026Q4净利润(亿元)',
    )
    actual_pattern = re.compile(
        r'([-+]?\d+(?:\.\d+)?)（(?:上个Q预测|锁定预测|未锁定预测)'
        r'([-+]?\d+(?:\.\d+)?)）')
    for row in target_by_code.values():
        l1 = str(target.cell(row, headers['申万一级']).value or '').strip()
        l2 = str(target.cell(row, headers['申万二级']).value or '').strip()
        expected_fill = mapping['industry_pairs'][l1 + '\t' + l2]['b_color'][1:].upper()
        for header in ('名称', '申万一级', '申万二级', '代码'):
            if fill(target.cell(row, headers[header])) != expected_fill:
                industry_fill_errors.append((row, header))
        for header in pe_headers:
            cell = target.cell(row, headers[header])
            value = number(cell.value)
            expected = ('00B050' if value is not None and value < 30 else
                        'C00000' if value is not None and value > 50 else '000000')
            if color(cell) != expected:
                font_rule_errors.append((row, header, expected, color(cell)))
        for header in yoy_headers:
            cell = target.cell(row, headers[header])
            value = number(cell.value)
            expected = ('C00000' if value is not None and value <= 0 else
                        '00B050' if value is not None and value >= .30 else
                        '0070C0' if value is not None and value >= .10 else '000000')
            if color(cell) != expected:
                font_rule_errors.append((row, header, expected, color(cell)))
        for header in actual_headers:
            cell = target.cell(row, headers[header])
            match = actual_pattern.search(str(cell.value or ''))
            expected = '000000'
            if match:
                actual, forecast = map(float, match.groups())
                ratio = ((actual - forecast) / abs(forecast) if forecast else
                         (0 if abs(actual) <= .05 else math.copysign(1, actual)))
                expected = ('0070C0' if abs(ratio) <= .10 else
                            '00B050' if ratio > .10 else 'C00000')
            if color(cell) != expected:
                font_rule_errors.append((row, header, expected, color(cell)))

    expected_formats = {
        '市值(亿元)': '#,##0',
        '未来三年增速均值': '0.0%',
        '2025净利润同比%(实际/预测)': '0.0%',
        '2026净利润同比%(预测)': '0.0%',
        '当日涨幅%': '0.0%', '10日涨幅%': '0.0%', '30日涨幅%': '0.0%',
        '60日涨幅%': '0.0%', '90日涨幅%': '0.0%', '250日涨幅%': '0.0%',
        '360日涨幅%': '0.0%', '360日内峰值距今涨幅%': '0.0%',
        '720日涨幅%': '0.0%', '1080日涨幅%': '0.0%',
        '26年预测变化率': '0.00%', '27年预测变化率': '0.00%',
        '28年预测变化率': '0.00%', '股息率%': '0.00%',
        '上财年全年股息率（25年·年日均市值）': '0.00%',
        '本财年预估全年股息率（26年）': '0.00%',
    }
    format_errors = []
    for header, expected in expected_formats.items():
        for row in target_by_code.values():
            actual = target.cell(row, headers[header]).number_format
            if actual != expected:
                format_errors.append((row, header, expected, actual))

    result = {
        'rows': len(target_by_code),
        'added': sorted(set(target_by_code) - set(source_by_code)),
        'removed': sorted(set(source_by_code) - set(target_by_code)),
        'latest_dates': sorted({str(target.cell(row, headers['最新交易日']).value)
                                for row in target_by_code.values()}),
        'note_style_or_value_mismatches': len(note_mismatches),
        'composition_mismatches': len(composition_mismatches),
        'industry_fill_errors': len(industry_fill_errors),
        'font_rule_errors': len(font_rule_errors),
        'number_format_errors': len(format_errors),
        'freeze': str(target.freeze_panes),
        'conditional_formats_exported': len(target.conditional_formatting),
    }
    assert result['rows'] == 462
    assert result['added'] == ['603225.SH', '603515.SH']
    assert result['removed'] == []
    assert result['latest_dates'] == ['2026-10-29']
    assert not note_mismatches and not composition_mismatches
    assert not industry_fill_errors and not font_rule_errors and not format_errors
    assert result['freeze'] == 'E2'
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
