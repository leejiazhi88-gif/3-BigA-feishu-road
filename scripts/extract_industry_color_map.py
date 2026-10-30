import json
import sys
from pathlib import Path

import openpyxl


def main(source, output, source_sheet):
    sheet = openpyxl.load_workbook(source, data_only=False)[source_sheet]
    headers = {cell.value: cell.column for cell in sheet[1]}
    l1_column, l2_column = headers['申万一级'], headers['申万二级']
    mappings = {'reference_sheet': source_sheet, 'industry_pairs': {}}
    for row in range(2, sheet.max_row + 1):
        l1 = str(sheet.cell(row, l1_column).value or '').strip()
        l2 = str(sheet.cell(row, l2_column).value or '').strip()
        if not l1 or not l2:
            continue
        colors = []
        for column in (l1_column, l2_column):
            cell = sheet.cell(row, column)
            color = cell.fill.fgColor.rgb
            assert cell.fill.fill_type == 'solid' and isinstance(color, str)
            colors.append('#' + color[-6:])
        key = l1 + '\t' + l2
        value = {'sw_l1': l1, 'sw_l2': l2, 'b_color': colors[0], 'c_color': colors[1]}
        previous = mappings['industry_pairs'].setdefault(key, value)
        assert previous == value, (key, previous, value)
    assert len(mappings['industry_pairs']) == 131
    Path(output).write_text(json.dumps(mappings, ensure_ascii=False, indent=2))
    print(json.dumps({
        'reference_sheet': source_sheet,
        'industry_pairs': len(mappings['industry_pairs']),
        'sw_l1': len({x['sw_l1'] for x in mappings['industry_pairs'].values()}),
        'sw_l2': len({x['sw_l2'] for x in mappings['industry_pairs'].values()}),
        'colors': sorted({
            color for item in mappings['industry_pairs'].values()
            for color in (item['b_color'], item['c_color'])
        }),
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], sys.argv[3])
